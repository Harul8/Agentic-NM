#!/usr/bin/env python3
"""
Read Excel (.xlsx) and Word (.docx) files and display in console.
Usage:
    python scripts/read_docs.py file.xlsx
    python scripts/read_docs.py file.docx
"""

import sys
import os
from pathlib import Path

def read_excel(filepath):
    """Read and display Excel file."""
    try:
        import pandas as pd
        df = pd.read_excel(filepath)
        print(f"\n{'='*80}")
        print(f"Excel File: {os.path.basename(filepath)}")
        print(f"{'='*80}\n")
        print(df.to_string(index=False))
        print(f"\n{'='*80}")
        print(f"Shape: {df.shape[0]} rows × {df.shape[1]} columns")
        print(f"{'='*80}\n")
    except Exception as e:
        print(f"Error reading Excel: {e}", file=sys.stderr)
        sys.exit(1)

def read_word(filepath):
    """Read and display Word file."""
    try:
        from docx import Document
        doc = Document(filepath)
        print(f"\n{'='*80}")
        print(f"Word Document: {os.path.basename(filepath)}")
        print(f"{'='*80}\n")
        for i, para in enumerate(doc.paragraphs, 1):
            text = para.text.strip()
            if text:
                print(f"{i}. {text}")
        print(f"\n{'='*80}")
        print(f"Total paragraphs: {len([p for p in doc.paragraphs if p.text.strip()])}")
        print(f"{'='*80}\n")
    except Exception as e:
        print(f"Error reading Word: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python scripts/read_docs.py <file.xlsx|file.docx>")
        sys.exit(1)
    
    filepath = sys.argv[1]
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}", file=sys.stderr)
        sys.exit(1)
    
    ext = Path(filepath).suffix.lower()
    if ext == ".xlsx":
        read_excel(filepath)
    elif ext == ".docx":
        read_word(filepath)
    else:
        print(f"Unsupported file type: {ext}. Use .xlsx or .docx", file=sys.stderr)
        sys.exit(1)
